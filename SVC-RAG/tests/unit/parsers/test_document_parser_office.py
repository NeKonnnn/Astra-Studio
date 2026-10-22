"""services/document_parser: docx/xlsx/xls/pptx/ppt, маршрутизация parse_document, откат при отсутствии библиотек."""

from __future__ import annotations

import importlib
import logging
import sys
from io import BytesIO
from types import SimpleNamespace

import pytest
from lxml import etree

from app.services import document_parser as dp

from . import pptx_factory as pxf
from . import ppt_factory as pf

pytestmark = pytest.mark.unit

docx = pytest.importorskip("docx")
openpyxl = pytest.importorskip("openpyxl")
PIL_Image = pytest.importorskip("PIL.Image")

W = dp._W_NS
W_URI = W.strip("{}")


def png_bytes(size=(40, 40), color=(10, 120, 200), fmt="PNG"):
    buf = BytesIO()
    PIL_Image.new("RGB", size, color).save(buf, fmt)
    return buf.getvalue()


def el(xml: str):
    return etree.fromstring(f'<root xmlns:w="{W_URI}">{xml}</root>')[0]


@pytest.fixture
def office_ocr(monkeypatch):
    """Подмена OCR картинок офисных файлов; пишет вызовы."""
    state = SimpleNamespace(text="", calls=[], per_name={})

    async def office(blobs, filename):
        state.calls.append((list(blobs), filename))
        return state.text

    monkeypatch.setattr(dp, "_ocr_office_images", office)
    return state


# ───────────────────────────── docx: XML-хелперы ─────────────────────────────


class TestDocxHelpers:
    @pytest.mark.parametrize("xml,expected", [
        ("<w:tc/>", 1),
        ("<w:tc><w:tcPr/></w:tc>", 1),
        ('<w:tc><w:tcPr><w:gridSpan w:val="3"/></w:tcPr></w:tc>', 3),
        ('<w:tc><w:tcPr><w:gridSpan w:val="0"/></w:tcPr></w:tc>', 1),
        ('<w:tc><w:tcPr><w:gridSpan w:val="x"/></w:tcPr></w:tc>', 1),
        ("<w:tc><w:tcPr><w:gridSpan/></w:tcPr></w:tc>", 1),
    ])
    def test_grid_span(self, xml, expected):
        assert dp._docx_tc_grid_span(el(xml)) == expected

    @pytest.mark.parametrize("xml,expected", [
        ("<w:tc/>", None),
        ("<w:tc><w:tcPr/></w:tc>", None),
        ('<w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr></w:tc>', "restart"),
        ('<w:tc><w:tcPr><w:vMerge w:val=" RESTART "/></w:tcPr></w:tc>', "restart"),
        ("<w:tc><w:tcPr><w:vMerge/></w:tcPr></w:tc>", "continue"),
        ('<w:tc><w:tcPr><w:vMerge w:val="continue"/></w:tcPr></w:tc>', "continue"),
    ])
    def test_vmerge(self, xml, expected):
        assert dp._docx_tc_vmerge(el(xml)) == expected

    def test_cell_text_collects_runs_tabs_and_ignores_deleted_text_and_comments(self):
        tc = el("<w:tc><w:p><w:r><w:t>Раз</w:t><w:tab/><w:t>два</w:t></w:r><!-- комментарий -->"
                "<w:r><w:br/><w:t>три</w:t></w:r><w:r><w:delText>удалено</w:delText></w:r>"
                "<w:ins><w:r><w:t>добавлено</w:t></w:r></w:ins><w:r><w:cr/><w:t/></w:r></w:p></w:tc>")

        assert dp._docx_tc_text(tc) == "Раз два тридобавлено"  # между w:r слов нет разделителя - так и в Word

    def test_table_rows_from_plain_element_and_from_python_docx_table(self):
        plain = el("<w:tbl><w:tr/><w:tr/></w:tbl>")
        assert len(dp._docx_table_rows(plain)) == 2

        document = docx.Document()
        table = document.add_table(rows=3, cols=2)
        assert len(dp._docx_table_rows(table)) == 3


class TestDocxCellParagraphs:
    def _table_docx(self):
        document = docx.Document()
        cell = document.add_table(rows=1, cols=2).cell(0, 0)
        cell.text = "Иван"
        cell.add_paragraph("Петров")
        document.tables[0].cell(0, 1).text = "x"
        buf = BytesIO()
        document.save(buf)
        return buf.getvalue()

    def test_paragraphs_inside_a_cell_must_be_separated(self):
        assert "| Иван Петров | x |" in dp.extract_text_from_docx(self._table_docx())


class TestDocxExtraction:
    def test_not_available(self, monkeypatch):
        monkeypatch.setattr(dp, "DOCX_AVAILABLE", False)

        with pytest.raises(RuntimeError, match="python-docx не установлен"):
            dp.extract_text_from_docx(b"x")

    def test_broken_table_is_skipped_and_logged(self, monkeypatch, caplog):
        document = docx.Document()
        document.add_paragraph("до таблицы")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "ячейка"
        document.add_paragraph("после таблицы")
        buf = BytesIO()
        document.save(buf)

        def boom(table):
            raise ValueError("кривая таблица")

        monkeypatch.setattr(dp, "_docx_table_to_markdown", boom)

        with caplog.at_level(logging.WARNING, logger=dp.logger.name):
            text = dp.extract_text_from_docx(buf.getvalue())

        assert "до таблицы" in text and "после таблицы" in text and "ячейка" not in text
        assert any("таблица #1 не разобрана (ValueError: кривая таблица)" in r.getMessage() for r in caplog.records)

    def test_image_blobs_of_document_with_picture(self):
        document = docx.Document()
        document.add_picture(BytesIO(png_bytes()))
        buf = BytesIO()
        document.save(buf)

        blobs = dp._docx_image_blobs(buf.getvalue())

        assert len(blobs) == 1 and blobs[0][:8] == b"\x89PNG\r\n\x1a\n"

    def test_image_blob_failure_of_single_relationship_is_skipped(self, monkeypatch):
        class BadRel:
            reltype = "http://x/relationships/image"

            @property
            def target_part(self):
                raise KeyError("нет части")

        fake_doc = SimpleNamespace(part=SimpleNamespace(rels={"a": BadRel(), "b": SimpleNamespace(reltype="http://x/hyperlink")}))
        monkeypatch.setattr(dp.docx, "Document", lambda f: fake_doc)

        assert dp._docx_image_blobs(b"x") == []

    def test_image_blobs_of_broken_file(self, caplog):
        with caplog.at_level(logging.WARNING, logger=dp.logger.name):
            assert dp._docx_image_blobs(b"not a docx") == []

        assert any("не удалось извлечь изображения" in r.getMessage() for r in caplog.records)


# ───────────────────────────── xlsx / xls ─────────────────────────────


def xlsx_with_image():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws["A1"] = "значение"
    from openpyxl.drawing.image import Image as XlImage

    ws.add_image(XlImage(BytesIO(png_bytes((20, 20)))), "C3")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestXlsx:
    def test_not_available(self, monkeypatch):
        monkeypatch.setattr(dp, "OPENPYXL_AVAILABLE", False)

        with pytest.raises(RuntimeError, match="openpyxl не установлен"):
            dp.extract_text_from_xlsx(b"x")

    def test_picture_blobs_are_extracted(self):
        blobs = dp._xlsx_image_blobs(xlsx_with_image())

        assert len(blobs) == 1 and blobs[0][:4] == b"\x89PNG"

    def test_picture_without_data_or_failing_data_is_skipped(self, monkeypatch):
        class Img:
            def __init__(self, data=None, error=None):
                self._d, self._e = data, error

            def _data(self):
                if self._e:
                    raise self._e
                return self._d

        ws = SimpleNamespace(_images=[Img(None), Img(error=OSError("x")), Img(b"ok"), SimpleNamespace(_data=None)])
        monkeypatch.setattr(dp.openpyxl, "load_workbook", lambda f: SimpleNamespace(worksheets=[ws, SimpleNamespace()]))

        assert dp._xlsx_image_blobs(b"x") == [b"ok"]

    def test_broken_workbook_must_give_empty_list(self):
        assert dp._xlsx_image_blobs(b"not an xlsx") == []


class FakeCell:
    def __init__(self, value):
        self.value = value


class FakeSheet:
    def __init__(self, name, rows):
        self.name, self._rows = name, rows
        self.nrows = len(rows)

    def row(self, i):
        return [FakeCell(v) for v in self._rows[i]]


class TestXls:
    def test_not_available(self, monkeypatch):
        monkeypatch.setattr(dp, "XLRD_AVAILABLE", False)

        with pytest.raises(RuntimeError, match="xlrd не установлен"):
            dp.extract_text_from_xls_xlrd(b"x")

    def test_sheets_rows_and_empty_cells(self, monkeypatch):
        sheets = [FakeSheet("Данные", [["a", "", None, 3.0], ["", None], ["x", "y"]]), FakeSheet("Пусто", [])]
        book = SimpleNamespace(nsheets=2, sheet_by_index=lambda i: sheets[i])
        seen = {}

        def open_workbook(file_contents):
            seen["data"] = file_contents
            return book

        monkeypatch.setattr(dp, "xlrd", SimpleNamespace(open_workbook=open_workbook), raising=False)
        monkeypatch.setattr(dp, "XLRD_AVAILABLE", True)

        text = dp.extract_text_from_xls_xlrd(b"XLSDATA")

        assert text == "Лист: Данные\na\t3.0\nx\ty\nЛист: Пусто" and seen["data"] == b"XLSDATA"


# ───────────────────────────── pptx ─────────────────────────────


class FakeCellP:
    def __init__(self, text):
        self.text = text


class FakeRow:
    def __init__(self, cells):
        self.cells = cells


class FakeTable:
    def __init__(self, rows):
        self.rows = rows


def table_of(*rows):
    return FakeTable([FakeRow([FakeCellP(c) for c in row]) for row in rows])


class TestPptxTable:
    def test_header_separator_and_whitespace(self):
        t = table_of(["Имя", " Воз   раст "], ["Анна", "30"])

        assert dp._pptx_table_to_markdown(t) == "| Имя | Воз раст |\n| --- | --- |\n| Анна | 30 |"

    def test_merged_cell_object_repeated_across_columns_is_taken_once(self):
        merged = FakeCellP("Шапка")
        t = FakeTable([FakeRow([merged, merged, FakeCellP("C")]), FakeRow([FakeCellP("1"), FakeCellP("2"), FakeCellP("3")])])

        assert dp._pptx_table_to_markdown(t) == "| Шапка | C |\n| --- | --- |\n| 1 | 2 | 3 |"

    def test_empty_rows_skipped_and_none_text(self):
        t = FakeTable([FakeRow([FakeCellP(None)]), FakeRow([FakeCellP("")])])

        assert dp._pptx_table_to_markdown(t) == ""

    def test_empty_first_row_loses_header_separator_known_issue(self):
        # RAG-23 (вариант python-pptx): разделитель только при r_idx == 0, в docx это исправлено счётчиком emitted.
        t = table_of(["", ""], ["Имя", "Возраст"], ["Анна", "30"])

        assert dp._pptx_table_to_markdown(t) == "| Имя | Возраст |\n| Анна | 30 |"

    @pytest.mark.xfail(strict=True, reason="BUG RAG-23: пустая первая строка таблицы pptx - нет разделителя шапки |---|")
    def test_empty_first_row_must_keep_separator(self):
        t = table_of(["", ""], ["Имя", "Возраст"], ["Анна", "30"])

        assert dp._pptx_table_to_markdown(t) == "| Имя | Возраст |\n| --- | --- |\n| Анна | 30 |"

    def test_single_row_has_no_separator(self):
        assert dp._pptx_table_to_markdown(table_of(["a", "b"])) == "| a | b |"


@pytest.fixture
def ocr_names(monkeypatch):
    """OCR картинок слайдов: имя -> текст; пишет, что и под каким именем отправили."""
    state = SimpleNamespace(answers={}, sent=[], error_for=set())

    async def ocr(blob, name):
        state.sent.append((name, blob))
        if name in state.error_for:
            raise RuntimeError("OCR упал")
        return state.answers.get(name, "")

    monkeypatch.setattr(dp, "_ocr_image_bytes", ocr)
    return state


class TestPptxViaXml:
    async def test_xml_fallback_when_library_is_missing(self, monkeypatch, ocr_names, caplog):
        monkeypatch.setattr(dp, "PPTX_AVAILABLE", False)
        ocr_names.answers = {"slide1_img1.png": "текст картинки"}
        data = pxf.build_pptx(
            [pxf.slide_xml(pxf.sp("Заголовок"), pxf.pic("rId1"), pxf.table([["a", "b"], ["1", "2"]]))],
            {1: [("rId1", pxf.T_IMAGE, "../media/i.png"), ("rId2", pxf.T_NOTES, "../notesSlides/n.xml")]},
            {"ppt/media/i.png": pxf.PNG_1x1, "ppt/notesSlides/n.xml": pxf.notes_xml("заметка")},
        )

        with caplog.at_level(logging.INFO, logger=dp.logger.name):
            text = await dp.extract_text_from_pptx(data)

        assert text == ("[Слайд 1]\n\nЗаголовок\n\n[Текст с изображения (OCR)]: текст картинки\n\n"
                        "| a | b |\n| --- | --- |\n| 1 | 2 |\n\n[Заметки докладчика]: заметка")
        assert any("python-pptx не установлен" in r.getMessage() for r in caplog.records)

    async def test_image_limit_comes_from_env(self, monkeypatch, ocr_names):
        monkeypatch.setattr(dp, "PPTX_AVAILABLE", False)
        monkeypatch.setenv("RAG_MAX_OCR_IMAGES_PER_DOC", "1")
        rels = [("rId1", pxf.T_IMAGE, "../media/1.png"), ("rId2", pxf.T_IMAGE, "../media/2.png")]
        data = pxf.build_pptx([pxf.slide_xml(pxf.pic("rId1"), pxf.pic("rId2"))], {1: rels},
                              {"ppt/media/1.png": b"1", "ppt/media/2.png": b"2"})
        ocr_names.answers = {"slide1_img1.png": "первая", "slide1_img2.png": "вторая"}

        text = await dp.extract_text_from_pptx(data)

        assert "первая" in text and "вторая" not in text


class FakeShape(SimpleNamespace):
    pass


class RaisingAttr:
    """Объект, любое чтение has_* которого падает - проверяем, что фигура не роняет разбор."""

    def __getattr__(self, name):
        raise RuntimeError("атрибут недоступен")


def install_fake_pptx(monkeypatch, slides):
    prs = SimpleNamespace(slides=slides)
    monkeypatch.setitem(sys.modules, "pptx", SimpleNamespace(Presentation=lambda f: prs))
    monkeypatch.setattr(dp, "PPTX_AVAILABLE", True)
    return prs


def slide_of(shapes, notes=None, has_notes=None):
    notes_slide = SimpleNamespace(notes_text_frame=SimpleNamespace(text=notes))
    return SimpleNamespace(shapes=shapes, has_notes_slide=(notes is not None) if has_notes is None else has_notes,
                           notes_slide=notes_slide)


class TestPptxViaLibrary:
    async def test_shapes_tables_pictures_and_notes(self, monkeypatch, ocr_names):
        blob = png_bytes()
        shapes = [
            FakeShape(has_table=False, has_text_frame=True, text_frame=SimpleNamespace(text="  Заголовок  "), shape_type=17),
            FakeShape(has_table=True, table=table_of(["x", "y"], ["1", "2"])),
            FakeShape(has_table=True, table=table_of([""])),  # пустая таблица: в слайд ничего
            FakeShape(has_table=False, has_text_frame=False, shape_type=13, image=SimpleNamespace(blob=blob)),
            FakeShape(has_table=False, has_text_frame=True, text_frame=SimpleNamespace(text=None), shape_type=17),
            FakeShape(has_table=False, has_text_frame=False, shape_type=13, image=SimpleNamespace(blob=b"")),
        ]
        install_fake_pptx(monkeypatch, [slide_of(shapes, notes="  заметка ")])
        ocr_names.answers = {"slide1_s3.png": "распознано"}

        text = await dp.extract_text_from_pptx(b"ignored")

        assert text == ("[Слайд 1]\n\nЗаголовок\n\n| x | y |\n| --- | --- |\n| 1 | 2 |\n\n"
                        "[Текст с изображения (OCR)]: распознано\n\n[Заметки докладчика]: заметка")
        assert ocr_names.sent == [("slide1_s3.png", blob)]

    async def test_broken_shapes_and_notes_do_not_break_the_slide(self, monkeypatch, ocr_names):
        shapes = [RaisingAttr(), FakeShape(has_table=False, has_text_frame=True, text_frame=SimpleNamespace(text="остался"), shape_type=1)]
        bad_notes = SimpleNamespace(shapes=[], has_notes_slide=True, notes_slide=RaisingAttr())
        install_fake_pptx(monkeypatch, [slide_of(shapes), bad_notes])

        text = await dp.extract_text_from_pptx(b"x")

        assert text == "[Слайд 1]\n\nостался\n\n[Слайд 2]"

    async def test_picture_limit_is_logged_and_extra_pictures_are_not_ocr_ed(self, monkeypatch, ocr_names, caplog):
        monkeypatch.setenv("RAG_MAX_OCR_IMAGES_PER_DOC", "1")
        pics = [FakeShape(has_table=False, has_text_frame=False, shape_type=13, image=SimpleNamespace(blob=png_bytes()))] * 2
        install_fake_pptx(monkeypatch, [slide_of(pics)])

        with caplog.at_level(logging.WARNING, logger=dp.logger.name):
            await dp.extract_text_from_pptx(b"x")

        assert len(ocr_names.sent) == 1
        assert any("лимит картинок 1 превышен" in r.getMessage() for r in caplog.records)

    async def test_library_errors_propagate(self, monkeypatch):
        def boom(f):
            raise ValueError("не pptx")

        monkeypatch.setitem(sys.modules, "pptx", SimpleNamespace(Presentation=boom))
        monkeypatch.setattr(dp, "PPTX_AVAILABLE", True)

        with pytest.raises(ValueError):
            await dp.extract_text_from_pptx(b"x")


class TestSlidesText:
    async def test_non_png_images_are_converted_for_ocr(self, ocr_names):
        bmp = png_bytes(fmt="BMP")
        ocr_names.answers = {"pic.bmp": "с bmp"}

        text = await dp._slides_text([[("img", "pic.bmp")]], [("pic.bmp", bmp)], [""], label="PPTX")

        assert "[Текст с изображения (OCR)]: с bmp" not in text  # ответ привязан к имени «pic.png», а не «pic.bmp»
        name, sent = ocr_names.sent[0]
        assert name == "pic.png" and sent[:8] == b"\x89PNG\r\n\x1a\n"

    async def test_png_and_jpeg_are_sent_as_is(self, ocr_names):
        ocr_names.answers = {"a.png": "A", "b.jpeg": "B"}

        text = await dp._slides_text([[("img", "a.png"), ("img", "b.jpeg")]], [("a.png", b"PNGDATA"), ("b.jpeg", b"JPG")], [""], label="X")

        assert ocr_names.sent == [("a.png", b"PNGDATA"), ("b.jpeg", b"JPG")]
        assert "A" in text and "B" in text

    async def test_unreadable_image_is_skipped_with_warning(self, ocr_names, caplog):
        with caplog.at_level(logging.WARNING, logger=dp.logger.name):
            text = await dp._slides_text([[("img", "bad.bmp")]], [("bad.bmp", b"not an image")], [""], label="PPT")

        assert ocr_names.sent == [] and text == "[Слайд 1]"
        assert any("картинка bad.bmp не открылась" in r.getMessage() for r in caplog.records)

    async def test_ocr_failure_is_skipped(self, ocr_names, caplog):
        ocr_names.error_for = {"a.png"}

        with caplog.at_level(logging.WARNING, logger=dp.logger.name):
            text = await dp._slides_text([[("img", "a.png")]], [("a.png", b"x")], [""], label="PPTX")

        assert text == "[Слайд 1]" and any("OCR картинки a.png пропущена" in r.getMessage() for r in caplog.records)

    async def test_without_pillow_everything_goes_as_is(self, monkeypatch, ocr_names):
        monkeypatch.setattr(dp, "PIL_AVAILABLE", False)
        ocr_names.answers = {"x.bmp": "bmp-текст"}

        text = await dp._slides_text([[("img", "x.bmp")]], [("x.bmp", b"raw")], [""], label="PPTX")

        assert ocr_names.sent == [("x.bmp", b"raw")] and "bmp-текст" in text


# ───────────────────────────── ppt ─────────────────────────────


class TestPptExtraction:
    @pytest.fixture
    def catppt(self, monkeypatch):
        state = SimpleNamespace(text="из catppt", calls=[], error=None)

        async def extract(data, name=""):
            state.calls.append((data, name))
            if state.error:
                raise state.error
            return state.text

        monkeypatch.setattr(dp, "extract_ppt_text", extract)
        return state

    def _ppt(self, *texts):
        return pf.build_ppt([pf.Slide(shapes=[pf.shape(pf.textbox(pf.text_chars(t)))]) for t in texts])

    async def test_own_parser_result_is_used(self, catppt):
        text = await dp.extract_text_from_ppt(self._ppt("Первый", "Второй"), "p.ppt")

        assert text == "[Слайд 1]\n\nПервый\n\n[Слайд 2]\n\nВторой" and catppt.calls == []

    async def test_ocr_text_of_pictures_is_inserted(self, monkeypatch, catppt):
        pics, offsets = pf.pictures_stream(pf.blip(0xF01E, 0x6E0, b"\x89PNG\r\n\x1a\nbody"))
        data = pf.build_ppt([pf.Slide(shapes=[pf.shape(pf.fopt(1))])], bstore_children=[pf.fbse(offsets[0])], pictures=pics)

        async def ocr(blob, name):
            return "надпись на картинке"

        monkeypatch.setattr(dp, "_ocr_image_bytes", ocr)

        text = await dp.extract_text_from_ppt(data, "p.ppt")

        assert text == "[Слайд 1]\n\n[Текст с изображения (OCR)]: надпись на картинке" and catppt.calls == []

    async def test_no_text_falls_back_to_catppt(self, catppt, caplog):
        data = pf.build_ppt([pf.Slide(shapes=[pf.shape(pf.textbox(pf.text_chars("*")))])])

        with caplog.at_level(logging.WARNING, logger=dp.logger.name):
            text = await dp.extract_text_from_ppt(data, "empty.ppt")

        assert text == "из catppt" and catppt.calls == [(data, "empty.ppt")]
        assert any("свой разбор - текста нет" in r.getMessage() for r in caplog.records)

    async def test_own_parser_failure_falls_back_to_catppt(self, catppt, caplog):
        with caplog.at_level(logging.WARNING, logger=dp.logger.name):
            text = await dp.extract_text_from_ppt(b"definitely not ole", "")

        assert text == "из catppt" and catppt.calls == [(b"definitely not ole", "?")]
        assert any("свой разбор не удался (PptError" in r.getMessage() for r in caplog.records)

    async def test_catppt_errors_propagate(self, catppt):
        catppt.error = RuntimeError("catppt не установлен")

        with pytest.raises(RuntimeError, match="catppt"):
            await dp.extract_text_from_ppt(b"junk", "x.ppt")


# ───────────────────────────── parse_document: маршрутизация ─────────────────────────────


@pytest.fixture
def no_mem_release(monkeypatch):
    calls = []
    monkeypatch.setattr(dp.mem_release, "release", lambda reason, **kw: calls.append(reason))
    return calls


class TestParseDocumentRouting:
    async def test_release_is_called_with_file_name_even_on_unknown_format(self, no_mem_release):
        assert await dp.parse_document(b"data", "x.unknown") is None
        assert await dp.parse_document(b"data", "") is None

        assert no_mem_release == ["разбор файла x.unknown", "разбор файла ?"]

    async def test_parser_exception_is_logged_and_gives_none(self, monkeypatch, no_mem_release, caplog):
        async def boom(data, name):
            raise RuntimeError("парсер упал")

        monkeypatch.setattr(dp, "_parse_document_by_ext", boom)

        with caplog.at_level(logging.ERROR, logger=dp.logger.name):
            assert await dp.parse_document(b"x", "f.txt") is None

        assert any("Не удалось разобрать файл «f.txt»" in r.getMessage() for r in caplog.records)

    async def test_docx_without_library_gives_none(self, monkeypatch):
        monkeypatch.setattr(dp, "DOCX_AVAILABLE", False)

        assert await dp._parse_document_by_ext(b"x", "a.docx") is None

    async def test_docx_with_ocr_text_only_document(self, monkeypatch, office_ocr):
        document = docx.Document()
        buf = BytesIO()
        document.save(buf)
        office_ocr.text = "только картинка"

        parsed = await dp._parse_document_by_ext(buf.getvalue(), "scan.docm")

        assert parsed["text"] == "только картинка" and parsed["confidence_info"]["file_type"] == "docx"

    async def test_pptx_ok_and_error(self, monkeypatch):
        async def ok(data):
            return "[Слайд 1]"

        monkeypatch.setattr(dp, "extract_text_from_pptx", ok)
        parsed = await dp._parse_document_by_ext(b"x", "s.pptm")
        assert parsed["text"] == "[Слайд 1]" and parsed["confidence_info"]["file_type"] == "pptx"

        async def boom(data):
            raise ValueError("битый")

        monkeypatch.setattr(dp, "extract_text_from_pptx", boom)
        assert await dp._parse_document_by_ext(b"x", "s.pptx") is None

    async def test_xlsx_branches(self, monkeypatch, office_ocr):
        monkeypatch.setattr(dp, "OPENPYXL_AVAILABLE", False)
        assert await dp._parse_document_by_ext(b"x", "t.xlsx") is None
        monkeypatch.setattr(dp, "OPENPYXL_AVAILABLE", True)

        assert await dp._parse_document_by_ext(b"not xlsx", "t.xlsx") is None

        office_ocr.text = "текст картинки"
        parsed = await dp._parse_document_by_ext(xlsx_with_image(), "t.xlsm")
        assert parsed["text"].startswith("Лист: Лист1\nзначение") and parsed["text"].endswith("текст картинки")
        assert len(office_ocr.calls[-1][0]) == 1 and office_ocr.calls[-1][1] == "t.xlsm"
        assert parsed["confidence_info"]["file_type"] == "excel"

    async def test_xlsx_with_empty_text_and_ocr(self, monkeypatch, office_ocr):
        monkeypatch.setattr(dp, "extract_text_from_xlsx", lambda data: "")
        monkeypatch.setattr(dp, "_xlsx_image_blobs", lambda data: [b"i"])
        office_ocr.text = "только OCR"

        parsed = await dp._parse_document_by_ext(b"x", "t.xlsx")

        assert parsed["text"] == "только OCR"

    async def test_xls_branches(self, monkeypatch):
        monkeypatch.setattr(dp, "XLRD_AVAILABLE", False)
        assert await dp._parse_document_by_ext(b"x", "t.xls") is None
        monkeypatch.setattr(dp, "XLRD_AVAILABLE", True)

        monkeypatch.setattr(dp, "extract_text_from_xls_xlrd", lambda data: "Лист: A")
        parsed = await dp._parse_document_by_ext(b"\xd0\xcf\x11\xe0", "t.xls")
        assert parsed["text"] == "Лист: A" and parsed["confidence_info"]["file_type"] == "excel"

        def boom(data):
            raise ValueError("xlrd")

        monkeypatch.setattr(dp, "extract_text_from_xls_xlrd", boom)
        assert await dp._parse_document_by_ext(b"\xd0\xcf\x11\xe0", "t.xls") is None

    async def test_doc_and_ppt_branches(self, monkeypatch):
        seen = []

        async def doc(data, name=""):
            seen.append(("doc", name))
            return "из doc"

        async def ppt(data, name=""):
            seen.append(("ppt", name))
            return "из ppt"

        monkeypatch.setattr(dp, "extract_doc_text", doc)
        monkeypatch.setattr(dp, "extract_text_from_ppt", ppt)

        d = await dp._parse_document_by_ext(b"\xd0\xcf", "Отчёт.DOC")
        p = await dp._parse_document_by_ext(b"\xd0\xcf", "slides.ppt")

        assert (d["text"], d["confidence_info"]["file_type"]) == ("из doc", "doc")
        assert (p["text"], p["confidence_info"]["file_type"]) == ("из ppt", "ppt")
        assert seen == [("doc", "Отчёт.DOC"), ("ppt", "slides.ppt")]

    async def test_doc_and_ppt_errors_give_none(self, monkeypatch):
        async def boom(data, name=""):
            raise RuntimeError("утилиты нет")

        monkeypatch.setattr(dp, "extract_doc_text", boom)
        monkeypatch.setattr(dp, "extract_text_from_ppt", boom)

        assert await dp._parse_document_by_ext(b"x", "a.doc") is None
        assert await dp._parse_document_by_ext(b"x", "a.ppt") is None

    @pytest.mark.parametrize("name", ["a.txt", "a.MD", "a.markdown", "a.log", "a.csv"])
    async def test_text_extensions(self, name):
        parsed = await dp._parse_document_by_ext("привет".encode("utf-8"), name)

        assert parsed["text"] == "привет" and parsed["confidence_info"]["file_type"] == "txt"

    async def test_rtf(self):
        parsed = await dp._parse_document_by_ext(rb"{\rtf1\ansi Hello}", "a.rtf")

        assert parsed["text"].strip() == "Hello" and parsed["confidence_info"]["file_type"] == "rtf"

    async def test_image_extensions_go_to_image_ocr(self, monkeypatch):
        seen = []

        async def image(data, name):
            seen.append(name)
            return {"text": "t", "file_type": "image"}

        monkeypatch.setattr(dp, "extract_text_from_image_bytes", image)

        for name in ("a.JPG", "a.jpeg", "a.png", "a.webp", "a.bmp"):
            assert (await dp._parse_document_by_ext(b"x", name))["file_type"] == "image"
        assert await dp._parse_document_by_ext(b"x", "") is None
        assert seen == ["a.JPG", "a.jpeg", "a.png", "a.webp", "a.bmp"]

    async def test_pdf_signature_beats_extension(self, monkeypatch):
        seen = []

        async def pdf(data, name):
            seen.append(name)
            return {"text": "pdf"}

        monkeypatch.setattr(dp, "extract_text_from_pdf_bytes", pdf)

        assert await dp._parse_document_by_ext(b"%PDF-1.7 ...", "scan.dat") == {"text": "pdf"}
        assert await dp._parse_document_by_ext(b"%PDF-1.7 ...", "") == {"text": "pdf"}
        assert await dp._parse_document_by_ext(b"x", ".pdf") == {"text": "pdf"}
        assert seen == ["scan.dat", "document.pdf", ".pdf"]


# ───────────────────────────── откат при отсутствии библиотек ─────────────────────────────


class TestOptionalImports:
    @pytest.fixture
    def reload_after(self, monkeypatch):
        yield
        monkeypatch.undo()
        importlib.reload(dp)

    def test_all_flags_false_when_libraries_cannot_be_imported(self, monkeypatch, reload_after):
        for name in ("docx", "fitz", "openpyxl", "xlrd", "pdf2image", "PIL", "pptx"):
            monkeypatch.setitem(sys.modules, name, None)

        importlib.reload(dp)

        assert (dp.DOCX_AVAILABLE, dp.PYMUPDF_AVAILABLE, dp.OPENPYXL_AVAILABLE, dp.XLRD_AVAILABLE) == (False,) * 4
        assert (dp.PDF2IMAGE_AVAILABLE, dp.PIL_AVAILABLE, dp.PPTX_AVAILABLE) == (False,) * 3

    def test_pptx_flag_true_when_library_is_importable(self, monkeypatch, reload_after):
        monkeypatch.setitem(sys.modules, "pptx", SimpleNamespace(Presentation=object))

        importlib.reload(dp)

        assert dp.PPTX_AVAILABLE is True

    def test_flags_recover_after_reload(self):
        # страховка к предыдущему тесту: после его finalizer модуль снова видит установленные библиотеки
        assert dp.DOCX_AVAILABLE and dp.PYMUPDF_AVAILABLE and dp.OPENPYXL_AVAILABLE and dp.PIL_AVAILABLE
