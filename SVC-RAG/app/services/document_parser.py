import asyncio
import logging
import os
import re
from io import BytesIO
from typing import Any, Dict, List, Optional

import httpx

from app.core.config import get_settings
from app.core.http_verify import resolve_httpx_verify

logger = logging.getLogger(__name__)

try:
    import docx

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import PyPDF2

    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    import fitz  # PyMuPDF

    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    import pdfplumber

    PDFPLUMBER_AVAILABLE = True
except ImportError:
    PDFPLUMBER_AVAILABLE = False

try:
    import openpyxl

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import xlrd

    XLRD_AVAILABLE = True
except ImportError:
    XLRD_AVAILABLE = False

try:
    from pdf2image import convert_from_bytes as pdf_convert_from_bytes

    PDF2IMAGE_AVAILABLE = True
except ImportError:
    PDF2IMAGE_AVAILABLE = False

try:
    from PIL import Image

    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False

try:
    from pptx import Presentation  # noqa: F401

    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


def _create_confidence_info_for_text(text: str, confidence_per_word: float, file_type: str) -> Dict[str, Any]:
    """Создаёт структуру confidence_info, совместимую с backend."""
    tokens = (text or "").split()
    words_with_confidence = [{"word": token, "confidence": float(confidence_per_word)} for token in tokens]
    avg_confidence = confidence_per_word
    return {
        "confidence": avg_confidence,
        "text_length": len(text or ""),
        "file_type": file_type,
        "words": words_with_confidence,
    }


async def _call_ocr_service(image_bytes: bytes, filename: str, languages: str = "ru,en") -> Dict[str, Any]:
    """Вызов OCR (Surya) по URL из config.yml"""
    settings = get_settings()
    ocr_url = settings.ocr.url.rstrip("/")
    timeout = settings.ocr.timeout

    mime = "image/jpeg"
    lower = filename.lower()
    if lower.endswith(".png"):
        mime = "image/png"
    elif lower.endswith(".bmp"):
        mime = "image/bmp"
    elif lower.endswith(".webp"):
        mime = "image/webp"

    files = {"file": (filename, BytesIO(image_bytes), mime)}
    data = {"languages": languages}

    try:
        req_timeout = httpx.Timeout(timeout, connect=10.0, read=timeout, write=10.0)
        async with httpx.AsyncClient(
            timeout=req_timeout, verify=resolve_httpx_verify()
        ) as client:
            resp = await client.post(
                f"{ocr_url}/v1/ocr", files=files, data=data, headers={"Accept": "application/json"}
            )
            resp.raise_for_status()
            return resp.json()
    except Exception as e:
        logger.error("Ошибка OCR при обращении к %s: %s", ocr_url, e)
        raise



_IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".webp", ".bmp")


def empty_document_index_error(
    parsed: Optional[Dict[str, Any]], filename: str = ""
) -> str:
    """Текст 422, когда парсер вернул пустой text (особенно для фото/OCR)."""
    conf: Dict[str, Any] = {}
    if isinstance(parsed, dict):
        raw = parsed.get("confidence_info")
        if isinstance(raw, dict):
            conf = raw
    err = conf.get("error")
    if err:
        return f"OCR изображения: {err}"
    ftype = ""
    if isinstance(parsed, dict):
        ftype = str(parsed.get("file_type") or conf.get("file_type") or "").lower()
    name = (filename or "").lower()
    is_image = ftype == "image" or any(name.endswith(ext) for ext in _IMAGE_EXTS)
    if is_image:
        return (
            "Не удалось распознать текст на изображении. "
            "Проверьте доступность SVC-OCR (ocr.url) и загрузите фото с читаемым текстом."
        )
    return "Документ пустой"


def _mupdf_store_shrink() -> None:
    """Опустошить кэш распакованных объектов MuPDF (общий на процесс)."""
    if not PYMUPDF_AVAILABLE:
        return
    try:
        tools = getattr(fitz, "TOOLS", None)
        if tools is None or not hasattr(tools, "store_shrink"):
            return
        before = int(getattr(tools, "store_size", 0) or 0)
        tools.store_shrink(100)
        after = int(getattr(tools, "store_size", 0) or 0)
        logger.debug(
            "MuPDF: кэш сброшен, %s -> %s байт (освободили %s)",
            before,
            after,
            before - after,
        )
    except Exception as e:
        logger.debug("MuPDF: сбросить кэш не вышло (%s)", e)


def _pdf_image_blobs(file_data: bytes) -> List[bytes]:
    """Встроенные изображения страниц PDF (для OCR) через PyMuPDF.

    Гибридный PDF: хороший текстовый слой + скриншоты/сканы фрагментов.
    Мелкие иконки отсекаются по площади, дедуп по xref.
    """
    if not PYMUPDF_AVAILABLE:
        return []
    min_area = int(os.getenv("RAG_OCR_IMAGE_MIN_AREA", "12000"))  # ~110x110
    max_imgs = int(os.getenv("RAG_MAX_OCR_IMAGES_PER_DOC", "50"))
    try:
        doc = fitz.open(stream=file_data, filetype="pdf")
        out: List[bytes] = []
        seen = set()
        for _i in range(doc.page_count):
            page = doc.load_page(_i)
            try:
                infos = page.get_images(full=True) or []
            except Exception:
                continue
            for info in infos:
                xref = int(info[0] if isinstance(info, (tuple, list)) else (info.get("xref") or 0))
                if not xref or xref in seen:
                    continue
                seen.add(xref)
                if isinstance(info, (tuple, list)):
                    w = int(info[2] or 0)
                    h = int(info[3] or 0)
                else:
                    w = int(info.get("width") or 0)
                    h = int(info.get("height") or 0)
                if w and h and w * h < min_area:
                    continue
                if len(out) >= max_imgs:
                    doc.close()
                    return out
                try:
                    im = doc.extract_image(xref)
                    blob = im.get("image") or b""
                    if blob:
                        out.append(blob)
                except Exception:
                    continue
        doc.close()
        return out
    except Exception as e:
        logger.warning("PDF: не удалось извлечь изображения: %s", e)
        return []


def _pptx_table_to_markdown(table) -> str:
    """Таблица pptx -> markdown-pipe строки."""
    lines = []
    for r_idx, row in enumerate(table.rows):
        cells = []
        prev_cell = None
        for cell in row.cells:
            if cell is prev_cell:
                continue
            prev_cell = cell
            cells.append(" ".join((cell.text or "").split()))
        if not any(cells):
            continue
        lines.append("| " + " | ".join(cells) + " |")
        if r_idx == 0 and len(table.rows) > 1:
            lines.append("|" + " --- |" * len(cells))
    return "\n".join(lines)


async def extract_text_from_pptx(file_data: bytes) -> str:
    """PPTX: текст/таблицы слайдов + OCR картинок через SVC-OCR (Surya)."""
    if not PPTX_AVAILABLE:
        raise RuntimeError("python-pptx не установлен")

    def _collect():
        from pptx import Presentation

        prs = Presentation(BytesIO(file_data))
        max_imgs = int(os.getenv("RAG_MAX_OCR_IMAGES_PER_DOC", "50"))
        per_slide: List[List[tuple]] = []
        img_items: List[tuple] = []
        for si, slide in enumerate(prs.slides, 1):
            items: List[tuple] = []
            for idx, shape in enumerate(slide.shapes):
                try:
                    if getattr(shape, "has_table", False):
                        md = _pptx_table_to_markdown(shape.table)
                        if md:
                            items.append(("text", md))
                        continue
                except Exception:
                    pass
                try:
                    if getattr(shape, "has_text_frame", False):
                        t = (shape.text_frame.text or "").strip()
                        if t:
                            items.append(("text", t))
                except Exception:
                    pass
                try:
                    if int(shape.shape_type) == 13:  # PICTURE
                        blob = shape.image.blob
                        if blob:
                            if len(img_items) < max_imgs:
                                name = f"slide{si}_s{idx}.png"
                                items.append(("img", name))
                                img_items.append((name, blob))
                            else:
                                logger.warning(
                                    "PPTX: лимит картинок %s превышен — дальше без OCR",
                                    max_imgs,
                                )
                except Exception:
                    pass
            per_slide.append(items)
        notes: List[str] = []
        for slide in prs.slides:
            try:
                notes.append(
                    (
                        slide.notes_slide.notes_text_frame.text
                        if slide.has_notes_slide
                        else ""
                    )
                    or ""
                )
            except Exception:
                notes.append("")
        return per_slide, img_items, notes

    per_slide, img_items, notes = await asyncio.to_thread(_collect)

    ocr_by_name: Dict[str, str] = {}
    for name, blob in img_items:
        try:
            res = await _call_ocr_service(blob, name, languages="ru,en")
            t = (res.get("text") or "").strip() if isinstance(res, dict) else ""
            if t:
                ocr_by_name[name] = t
        except Exception as e:
            logger.warning("PPTX: OCR картинки %s пропущена: %s", name, e)

    parts = []
    for si, items in enumerate(per_slide, 1):
        slide_parts = [f"[Слайд {si}]"]
        for kind, value in items:
            if kind == "text":
                slide_parts.append(value)
            else:
                t = (ocr_by_name.get(value) or "").strip()
                if t:
                    slide_parts.append(f"[Текст с изображения (OCR)]: {t}")
        nt = notes[si - 1].strip() if si - 1 < len(notes) else ""
        if nt:
            slide_parts.append(f"[Заметки докладчика]: {nt}")
        parts.append("\n\n".join(slide_parts))
    return "\n\n".join(parts)


def _docx_table_to_markdown(table) -> str:
    """Таблица docx -> markdown-pipe строки. Одна строка таблицы = одна строка
    текста: чанкер (_protect_tables) узнаёт такой блок и не режет его посередине.
    Merged-ячейки python-docx повторяет по разу на колонку - дедупим по XML-элементу.
    """
    lines = []
    for r_idx, row in enumerate(table.rows):
        cells = []
        prev_tc = None
        for cell in row.cells:
            if cell._tc is prev_tc:
                continue
            prev_tc = cell._tc
            cells.append(" ".join((cell.text or "").split()))
        if not any(cells):
            continue
        lines.append("| " + " | ".join(cells) + " |")
        if r_idx == 0 and len(table.rows) > 1:
            lines.append("|" + " --- |" * len(cells))
    return "\n".join(lines)

def extract_text_from_docx(file_data: bytes) -> str:
    if not DOCX_AVAILABLE:
        raise RuntimeError("python-docx не установлен")
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(BytesIO(file_data))
    parts = []
    # Обходим тело В ПОРЯДКЕ ДОКУМЕНТА: абзацы и таблицы вперемешку, как в файле.
    # Старый вариант (все абзацы, потом все таблицы поячеечно) разрушал структуру
    # и порождал сотни мелких бесконтекстных чанков.
    for child in doc.element.body.iterchildren():
        tag = child.tag
        if tag.endswith("}p"):
            parts.append(Paragraph(child, doc).text)
        elif tag.endswith("}tbl"):
            md = _docx_table_to_markdown(Table(child, doc))
            if md:
                # Пустые строки вокруг таблицы, чтобы сплиттер не клеил её к абзацу.
                parts.append("")
                parts.append(md)
                parts.append("")
    return "\n".join(parts)


def _docx_image_blobs(file_data: bytes) -> list:
    """Байты встроенных изображений .docx (для OCR)."""
    try:
        doc = docx.Document(BytesIO(file_data))
        out = []
        for rel in doc.part.rels.values():
            if "image" in (getattr(rel, "reltype", "") or ""):
                try:
                    out.append(rel.target_part.blob)
                except Exception:
                    continue
        return out
    except Exception as e:
        logger.warning("DOCX: не удалось извлечь изображения: %s", e)
        return []

def _xlsx_image_blobs(file_data: bytes) -> list:
    """Байты картинок, вставленных на листы .xlsx (для OCR)."""
    try:
        wb = openpyxl.load_workbook(BytesIO(file_data))
        out = []
        for ws in wb.worksheets:
            for img in getattr(ws, "_images", []) or []:
                try:
                    data = img._data() if callable(getattr(img, "_data", None)) else None
                    if data:
                        out.append(data)
                except Exception:
                    continue
        return out
    except Exception as e:
        logger.warning("XLSX: не удалось извлечь изображения: %s", e)
        return []

async def _ocr_office_images(blobs: list, filename: str) -> str:
    """OCR списка изображений офисного файла → склеенный текст (или '')."""
    if not blobs:
        return ""
    texts = []
    for i, blob in enumerate(blobs, 1):
        try:
            res = await _call_ocr_service(blob, f"{filename}.img{i}.png")
            t = (res.get("text") or "").strip() if isinstance(res, dict) else ""
            if t:
                texts.append(t)
        except Exception as e:
            logger.warning("Office OCR картинки пропущена (%s): %s", filename, e)
    if not texts:
        return ""
    return "[Текст из изображений документа (OCR)]:\n" + "\n".join(texts)


def _pdf_magic_bytes(data: bytes) -> bool:
    return bool(data and len(data) >= 5 and data[:5] == b"%PDF-")


def _normalize_extracted_text(text: str) -> str:
    """Нормализация извлечённого текста перед чанкингом/эмбеддингом."""
    t = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    # Мягкие переносы и разрывы слов на стыках строк из PDF
    t = t.replace("\u00ad", "")
    t = re.sub(r"(?<=\w)-\n(?=\w)", "", t)
    # Убираем «рваные» пробелы, но сохраняем абзацы
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def _looks_like_low_quality_pdf_text(text: str, n_pages: int) -> bool:
    """Эвристика плохого текстового слоя PDF (часто сканы/битые шрифты)."""
    t = (text or "").strip()
    if not t:
        return True
    pages = max(int(n_pages or 1), 1)
    min_chars_expected = max(120, pages * 80)
    if len(t) < min_chars_expected:
        return True
    letters = sum(1 for ch in t if ch.isalpha())
    digits = sum(1 for ch in t if ch.isdigit())
    spaces = sum(1 for ch in t if ch.isspace())
    total = max(len(t), 1)
    alpha_ratio = letters / total
    # Очень мало букв при длинном тексте → вероятный "мусорный" слой
    if alpha_ratio < 0.35 and (digits + spaces) / total > 0.45:
        return True
    # Слишком много одиночных токенов/обрывков
    toks = re.findall(r"\w+", t, re.UNICODE)
    if toks:
        short_ratio = sum(1 for w in toks if len(w) <= 2) / len(toks)
        if short_ratio > 0.72:
            return True
    return False


async def extract_text_from_pdf_bytes(file_data: bytes) -> Dict[str, Any]:
    """
    PyMuPDF + pdfplumber + PyPDF2 — выбирается самый длинный извлечённый текст;
    при пустом или слишком коротком слое относительно числа страниц — OCR (pdf2image + ocr-service).
    """
    logger.info("PDF: извлечение текста, размер=%s байт", len(file_data))
    candidates: List[tuple] = []
    n_pages = 0

    if PYMUPDF_AVAILABLE:
        try:
            doc = fitz.open(stream=file_data, filetype="pdf")
            if getattr(doc, "needs_pass", False):
                try:
                    doc.authenticate("")
                except Exception:
                    pass
            n_pages = doc.page_count
            parts: List[str] = []
            for i in range(n_pages):
                parts.append(doc.load_page(i).get_text() or "")
            doc.close()
            t = _normalize_extracted_text("\n".join(parts))
            if t.strip():
                candidates.append(("pymupdf", t))
                logger.info("PDF: PyMuPDF извлёк %s символов, страниц=%s", len(t), n_pages)
        except Exception as e:
            logger.warning("PDF: PyMuPDF ошибка: %s", e)

    text_pb_parts: List[str] = []
    if PDFPLUMBER_AVAILABLE:
        try:
            with pdfplumber.open(BytesIO(file_data)) as pdf:
                n_pages = max(n_pages, len(pdf.pages))
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    text_pb_parts.append(page_text)
            text_pb = _normalize_extracted_text("\n".join(text_pb_parts))
            if text_pb.strip():
                candidates.append(("pdfplumber", text_pb))
                logger.info("PDF: pdfplumber извлёк %s символов", len(text_pb))
        except Exception as e:
            logger.warning("PDF: pdfplumber ошибка: %s", e)

    if PYPDF2_AVAILABLE:
        try:
            reader = PyPDF2.PdfReader(BytesIO(file_data))
            n_pages = max(n_pages, len(reader.pages))
            text_p2_parts: List[str] = []
            for page in reader.pages:
                text_p2_parts.append(page.extract_text() or "")
            text_p2 = _normalize_extracted_text("\n".join(text_p2_parts))
            if text_p2.strip():
                candidates.append(("pypdf2", text_p2))
                logger.info("PDF: PyPDF2 извлёк %s символов", len(text_p2))
        except Exception as e2:
            logger.warning("PDF: PyPDF2 ошибка: %s", e2)

    text = ""
    confidence_scores: List[float] = []
    if candidates:
        label, text = max(candidates, key=lambda x: len(x[1].strip()))
        logger.info("PDF: для индекса выбран движок «%s» (%s символов)", label, len(text.strip()))

    weak_layer = _looks_like_low_quality_pdf_text(text, n_pages)

    need_ocr = (not text.strip() or weak_layer) and PDF2IMAGE_AVAILABLE and PIL_AVAILABLE
    if need_ocr:
        if weak_layer:
            logger.info(
                "PDF: текстовый слой низкого качества (%s симв., страниц=%s) — пробуем OCR",
                len(text.strip()),
                max(n_pages, 1),
            )
        else:
            logger.info("PDF: текстовый слой пуст — rasterize+OCR (poppler + ocr-service)")
        try:
            try:
                images = pdf_convert_from_bytes(file_data, dpi=300)
            except Exception as e_dpi:
                logger.warning("PDF: pdf2image dpi=300 не удалось (%s), пробуем dpi=150", e_dpi)
                images = pdf_convert_from_bytes(file_data, dpi=150)
            logger.info("PDF: растеризация в %s страниц для OCR", len(images))

            ocr_text = ""
            ocr_confidence_scores: List[float] = []
            for i, image in enumerate(images):
                buf = BytesIO()
                image.save(buf, format="PNG")
                buf.seek(0)
                page_image_data = buf.getvalue()
                result = await _call_ocr_service(page_image_data, f"page_{i+1}.png", languages="ru,en")
                if result.get("success"):
                    page_text = result.get("text", "")
                    ocr_text += f"\n--- Страница {i+1} ---\n{page_text}\n"
                    ocr_confidence_scores.append(float(result.get("confidence", 50.0) or 50.0))
                else:
                    logger.warning("PDF OCR: страница %s: %s", i + 1, result.get("error", "Unknown"))
                    ocr_confidence_scores.append(50.0)

            ocr_text = _normalize_extracted_text(ocr_text)
            if ocr_text.strip():
                if len(ocr_text.strip()) >= len(text.strip()):
                    text = ocr_text
                    confidence_scores = ocr_confidence_scores
                    logger.info("PDF: использован OCR-текст (%s символов)", len(text.strip()))
                else:
                    logger.info(
                        "PDF: OCR не дал больше текста (%s vs %s симв.) — оставляем текстовый слой",
                        len(ocr_text.strip()),
                        len(text.strip()),
                    )
            else:
                logger.warning("PDF: OCR не вернул текст")
        except Exception as e:
            logger.warning("PDF: OCR недоступен или сбой (poppler/pdf2image/ocr-service): %s", e)


    # Гибридный PDF: хороший текстовый слой + встроенные картинки (скриншоты).
    # Сканы целиком уже обработаны полным OCR выше — здесь только text+images.
    if text.strip() and not weak_layer:
        try:
            img_blobs = await asyncio.to_thread(_pdf_image_blobs, file_data)
            _mupdf_store_shrink()
            img_text = await _ocr_office_images(img_blobs, "document.pdf")
            if img_text:
                text = (text + "\n\n" + img_text) if text.strip() else img_text
                logger.info(
                    "PDF: гибридный — добавлен OCR-текст %s встроенных картинок",
                    len(img_blobs) if img_blobs else 0,
                )
        except Exception as e:
            logger.warning("PDF: OCR встроенных картинок не удался: %s", e)

    if text.strip() and not confidence_scores:
        confidence_scores = [95.0] * max(1, n_pages or 1)

    avg_confidence = sum(confidence_scores) / len(confidence_scores) if confidence_scores else 0.0
    confidence_per_word = 100.0 if avg_confidence > 90.0 else avg_confidence
    text = _normalize_extracted_text(text)
    confidence_info = _create_confidence_info_for_text(text or "", confidence_per_word, "pdf")
    confidence_info["pages_processed"] = len(confidence_scores) if confidence_scores else (n_pages or 0)

    if not (text or "").strip():
        logger.error(
            "PDF: итоговый текст пуст (страниц≈%s). Нужны рабочие PyMuPDF/pdfplumber, для сканов — poppler в образе и "
            "доступный ocr-service (SVC_OCR_URL).",
            n_pages,
        )

    return {
        "text": text or "",
        "confidence_info": confidence_info,
        "file_type": "pdf",
        "pages": int(confidence_info.get("pages_processed") or 0),
    }


def extract_text_from_xls_xlrd(file_data: bytes) -> str:
    """Старый формат Excel .xls (не .xlsx)."""
    if not XLRD_AVAILABLE:
        raise RuntimeError("xlrd не установлен")
    book = xlrd.open_workbook(file_contents=file_data)
    parts: List[str] = []
    for si in range(book.nsheets):
        sh = book.sheet_by_index(si)
        parts.append(f"Лист: {sh.name}")
        for ri in range(sh.nrows):
            row = sh.row(ri)
            vals = [str(c.value) for c in row if c.value not in ("", None)]
            if vals:
                parts.append("\t".join(vals))
    return "\n".join(parts)


def extract_text_from_xlsx(file_data: bytes) -> str:
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError("openpyxl не установлен")
    workbook = openpyxl.load_workbook(BytesIO(file_data), data_only=True)
    parts = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        parts.append(f"Лист: {sheet_name}")
        for row in sheet.iter_rows():
            row_vals = [str(c.value) for c in row if c.value is not None]
            if row_vals:
                parts.append("\t".join(row_vals))
    return "\n".join(parts)


def extract_text_from_txt(file_data: bytes) -> str:
    for enc in ("utf-8", "cp1251", "latin-1", "koi8-r"):
        try:
            return file_data.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_data.decode("utf-8", errors="replace")


def extract_text_from_rtf(file_data: bytes) -> str:
    """RTF -> плоский текст."""
    from striprtf.striprtf import rtf_to_text
    return rtf_to_text(file_data.decode("utf-8", errors="ignore"))


async def extract_text_from_image_bytes(file_data: bytes) -> Dict[str, Any]:
    """Извлечение текста из изображения (OCR). Вызов идёт в ocr-service (Surya)."""
    print(f"Извлекаем текст из изображения с помощью Surya OCR (размер: {len(file_data)} байт)")
    if not PIL_AVAILABLE:
        result_text = "[Изображение. Для распознавания текста требуется Pillow и доступ к ocr-service.]"
        return {
            "text": result_text,
            "confidence_info": _create_confidence_info_for_text(result_text, 0.0, "image"),
        }

    img = Image.open(BytesIO(file_data)).convert("RGB")
    # После convert() у Image.format обычно None; ниже всегда шлём PNG.
    filename = "image.png"

    print(f"DEBUG: Изображение открыто, размер: {img.size}")

    # Увеличиваем маленькие изображения, как в backend
    min_side = 1024
    w, h = img.size
    if max(w, h) < min_side and max(w, h) > 0:
        scale = min_side / max(w, h)
        new_w = max(1, int(round(w * scale)))
        new_h = max(1, int(round(h * scale)))
        img = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
        print(f"DEBUG: Изображение увеличено для OCR: {w}x{h} -> {new_w}x{new_h}")

    buf = BytesIO()
    img.save(buf, format="PNG")
    image_to_send = buf.getvalue()

    try:
        result = await _call_ocr_service(image_to_send, filename, languages="ru,en")
    except Exception as e:
        print(f"Ошибка OCR: {e}")
        return {
            "text": "",
            "confidence_info": {
                "confidence": 0.0,
                "text_length": 0,
                "file_type": "image",
                "ocr_available": False,
                "error": str(e),
                "words": [],
            },
        }

    if not result.get("success"):
        error_msg = result.get("error", "Неизвестная ошибка")
        print(f"Surya OCR вернул ошибку: {error_msg}")
        return {
            "text": "",
            "confidence_info": {
                "confidence": 0.0,
                "text_length": 0,
                "file_type": "image",
                "ocr_available": False,
                "error": error_msg,
                "words": [],
            },
        }

    text = result.get("text", "") or ""
    words = result.get("words", []) or []
    avg_confidence = float(result.get("confidence", 0.0) or 0.0)

    if not text.strip():
        print("Surya OCR не смог извлечь текст из изображения (текст пустой)")
        return {
            "text": "",
            "confidence_info": {
                "confidence": 0.0,
                "text_length": 0,
                "file_type": "image",
                "ocr_available": False,
                "words": [],
            },
        }

    print(
        f"Surya OCR успешно извлек {len(text)} символов, {len(words)} слов, средняя уверенность: {avg_confidence:.2f}%"
    )

    confidence_info = {
        "confidence": avg_confidence,
        "text_length": len(text),
        "file_type": "image",
        "ocr_available": True,
        "words": words,
    }
    return {"text": text, "confidence_info": confidence_info}


async def parse_document(file_data: bytes, filename: str) -> Optional[Dict[str, Any]]:
    """
    По расширению файла выбираем парсер и возвращаем структуру:
    {"text": str, "confidence_info": dict}
    """
    name = (filename or "").lower()
    # Корректно извлекаем только сам суффикс (".docx", ".pdf" и т.п.)
    dot = name.rfind(".")
    ext = name[dot:] if dot != -1 else ""
    if ext != ".pdf" and _pdf_magic_bytes(file_data):
        logger.info(
            "parse_document: файл «%s» без .pdf, но с сигнатурой PDF — парсим как PDF",
            filename or "?",
        )
        return await extract_text_from_pdf_bytes(file_data)
    if ext in (".docx", ".docm"):
        if not DOCX_AVAILABLE:
            return None
        text = extract_text_from_docx(file_data)
        img_text = await _ocr_office_images(
            _docx_image_blobs(file_data), filename or "document.docx"
        )
        if img_text:
            text = (text + "\n\n" + img_text) if text.strip() else img_text
        return {
            "text": text,
            "confidence_info": _create_confidence_info_for_text(text, 100.0, "docx"),
        }
    if ext == ".pdf":
        return await extract_text_from_pdf_bytes(file_data)
    if ext in (".pptx", ".pptm"):
        if not PPTX_AVAILABLE:
            logger.warning("Парсинг .pptx/.pptm: python-pptx не установлен")
            return None
        try:
            text = await extract_text_from_pptx(file_data)
        except Exception as e:
            logger.error("Ошибка парсинга .pptx/.pptm: %s", e)
            return None
        return {
            "text": text,
            "confidence_info": _create_confidence_info_for_text(text, 100.0, "pptx"),
        }
    if ext in (".xlsx", ".xlsm"):
        if not OPENPYXL_AVAILABLE:
            logger.warning("Парсинг .xlsx/.xlsm: openpyxl не установлен")
            return None
        try:
            text = extract_text_from_xlsx(file_data)
        except Exception as e:
            logger.error("Ошибка парсинга .xlsx/.xlsm: %s", e)
            return None
        img_text = await _ocr_office_images(
            _xlsx_image_blobs(file_data), filename or "table.xlsx"
        )
        if img_text:
            text = (text + "\n\n" + img_text) if text.strip() else img_text
        return {
            "text": text,
            "confidence_info": _create_confidence_info_for_text(text, 100.0, "excel"),
        }
    if ext == ".xls":
        if not XLRD_AVAILABLE:
            logger.warning("Парсинг .xls: xlrd не установлен")
            return None
        try:
            text = extract_text_from_xls_xlrd(file_data)
        except Exception as e:
            logger.error("Ошибка парсинга .xls: %s", e)
            return None
        return {
            "text": text,
            "confidence_info": _create_confidence_info_for_text(text, 100.0, "excel"),
        }
    if ext in (".txt", ".md", ".markdown", ".log"):
        text = extract_text_from_txt(file_data)
        return {
            "text": text,
            "confidence_info": _create_confidence_info_for_text(text, 100.0, "txt"),
        }
    if ext == ".rtf":
        text = extract_text_from_rtf(file_data)
        return {
            "text": text,
            "confidence_info": _create_confidence_info_for_text(text, 100.0, "rtf"),
        }
    if ext in (".jpg", ".jpeg", ".png", ".webp", ".bmp"):
        return await extract_text_from_image_bytes(file_data)
    return None
